"""Export extracted reports to a formatted Excel workbook (.xlsx).

Two sheets:

* **Records (flat)** — one row per patient, every field in its own column
  (the spreadsheet-native shape, with a frozen header row and auto-filter).
* **Field matrix** — the same data transposed (fields as rows grouped into
  labelled sections, one column per patient) for easy side-by-side reading.

Requires ``openpyxl`` (listed in requirements.txt).
"""

from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .flatten import flatten_report, column_order
from .html_report import _SECTIONS
from .schema import GeneticReport

# Palette (matches the HTML matrix)
_ACCENT = "0E7C86"
_ACCENT_WEAK = "E2F1F2"
_SECTION = "D7E9EB"
_ZEBRA = "F2F6F9"
_INK = "141A1F"

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor=_ACCENT)
_SECTION_FONT = Font(bold=True, color=_ACCENT, size=10)
_SECTION_FILL = PatternFill("solid", fgColor=_SECTION)
_LABEL_FONT = Font(bold=True, color=_INK, size=10)
_ZEBRA_FILL = PatternFill("solid", fgColor=_ZEBRA)
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")
_THIN = Side(style="thin", color="D5DEE4")
_BORDER = Border(bottom=_THIN, right=_THIN)

_LONG_FIELDS = {
    "test.method_summary", "clinical.free_text", "result_summary", "interpretation",
    "recommendations", "incidental_findings", "secondary_findings",
    "carriership_findings", "clinical.hpo_terms",
}


def _s(value) -> str:
    return "" if value is None else str(value)


def _flat_sheet(wb: Workbook, reports: List[GeneticReport]) -> None:
    ws = wb.active
    ws.title = "Records (flat)"
    rows = [flatten_report(r) for r in reports]
    cols = column_order(rows)

    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, start=2):
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=r_idx, column=c, value=_s(row.get(col)))
            cell.alignment = _WRAP if col in _LONG_FIELDS else _TOP
            if r_idx % 2 == 0:
                cell.fill = _ZEBRA_FILL

    # Column widths: wider for long free-text columns
    for c, col in enumerate(cols, start=1):
        width = 46 if col in _LONG_FIELDS else max(14, min(30, len(col) + 3))
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    ws.sheet_view.showGridLines = False


def _matrix_sheet(wb: Workbook, reports: List[GeneticReport]) -> None:
    ws = wb.create_sheet("Field matrix")
    flat = [flatten_report(r) for r in reports]

    # Header row: Field | patient columns
    hdr = ws.cell(row=1, column=1, value="Field")
    hdr.font = _HEADER_FONT
    hdr.fill = _HEADER_FILL
    for c, r in enumerate(reports, start=2):
        label = r.patient.full_name or "(unknown)"
        sub = f"no. {r.patient.patient_no or '?'}"
        if r.patients_in_source and r.patients_in_source > 1:
            sub += f"  ·  split {r.patient_index + 1}/{r.patients_in_source}"
        cell = ws.cell(row=1, column=c, value=f"{label}\n{sub}")
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ncols = len(reports) + 1
    row_i = 2
    for title, fields in _SECTIONS:
        sc = ws.cell(row=row_i, column=1, value=title.upper())
        sc.font = _SECTION_FONT
        sc.fill = _SECTION_FILL
        for c in range(2, ncols + 1):
            ws.cell(row=row_i, column=c).fill = _SECTION_FILL
        row_i += 1
        for label, key, is_long in fields:
            lc = ws.cell(row=row_i, column=1, value=label)
            lc.font = _LABEL_FONT
            lc.alignment = _TOP
            for c, row in enumerate(flat, start=2):
                cell = ws.cell(row=row_i, column=c, value=_s(row.get(key)))
                cell.alignment = _WRAP if is_long else _TOP
                cell.border = _BORDER
            row_i += 1

    ws.column_dimensions["A"].width = 26
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 44
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30


def _build_workbook(reports: List[GeneticReport]) -> Workbook:
    wb = Workbook()
    _flat_sheet(wb, reports)
    _matrix_sheet(wb, reports)
    return wb


def write_xlsx(reports: List[GeneticReport], path: str) -> None:
    _build_workbook(reports).save(path)


def xlsx_bytes(reports: List[GeneticReport]) -> bytes:
    """Return the workbook as bytes (for streaming from a web response)."""
    buf = io.BytesIO()
    _build_workbook(reports).save(buf)
    return buf.getvalue()


def write_xlsx_streaming(rows, columns, path: str, extra_headers=None) -> None:
    """Write a large flat sheet row-by-row using openpyxl write-only mode.

    ``rows`` is an iterable of dicts (e.g. from ``flatten_report``); nothing is
    held in memory beyond the current row, so this scales to tens of thousands of
    records. Suitable for the one-row-per-patient export at 70k scale (the
    transposed matrix sheet is intentionally omitted — it does not scale to that
    many patient columns).
    """
    from openpyxl import Workbook as _WB
    from openpyxl.cell import WriteOnlyCell

    cols = list(columns) + list(extra_headers or [])
    wb = _WB(write_only=True)
    ws = wb.create_sheet("Records (flat)")
    header = []
    for c in cols:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        header.append(cell)
    ws.append(header)
    for row in rows:
        ws.append(["" if row.get(c) is None else str(row.get(c)) for c in cols])
    wb.save(path)
